from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/39-notion-client-requirements.md'
OUTPUT = ROOT / 'output/docx/timefit-client-requirements-notion-guide.docx'
FONT = 'Nanum Gothic'
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)

def set_font(run, size=None, color=None, bold=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if size: run.font.size = Pt(size)
    if color: run.font.color.rgb = color
    if bold is not None: run.bold = bold

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr(); shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), fill); tc_pr.append(shade)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{side}'))
        if node is None: node = OxmlElement(f'w:{side}'); tcMar.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')

def clean(text):
    text = re.sub(r'\[([^]]+)\]\([^)]+\)', r'\1', text)
    return text.replace('**', '').replace('`', '').strip()

def add_text(p, text, size=11, color=None, bold=False):
    run = p.add_run(clean(text)); set_font(run, size, color, bold); return run

def add_table(doc, rows):
    cols = len(rows[0]); table = doc.add_table(rows=0, cols=cols); table.alignment = WD_TABLE_ALIGNMENT.LEFT; table.autofit = False
    widths = [Inches(6.5 / cols)] * cols
    for r, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = widths[i]; cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cells[i])
            if r == 0: set_cell_shading(cells[i], 'E8EEF5')
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0); add_text(p, value, 9.5, BLUE if r == 0 else None, r == 0)
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr(); tcW = tc_pr.find(qn('w:tcW')); tcW.set(qn('w:w'), str(int(9360 / cols))); tcW.set(qn('w:type'), 'dxa')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def build():
    doc = Document(); sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    normal = doc.styles['Normal']; normal.font.name = FONT; normal._element.rPr.rFonts.set(qn('w:eastAsia'), FONT); normal.font.size = Pt(11); normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [('Heading 1',16,BLUE,18,10), ('Heading 2',13,BLUE,14,7), ('Heading 3',12,DARK,10,5)]:
        style=doc.styles[name]; style.font.name=FONT; style._element.rPr.rFonts.set(qn('w:eastAsia'), FONT); style.font.size=Pt(size); style.font.color.rgb=color; style.paragraph_format.space_before=Pt(before); style.paragraph_format.space_after=Pt(after)
    header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT; add_text(header, 'TimeFit | 클라이언트 요구사항', 8.5, RGBColor(110,120,133))
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT; add_text(footer, 'TimeFit 내부 검토용', 8.5, RGBColor(110,120,133))
    lines = SOURCE.read_text(encoding='utf-8').splitlines(); table_rows=[]
    for raw in lines:
        text=raw.strip()
        if text.startswith('|'):
            cells=[x.strip() for x in text.strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+', c) for c in cells): table_rows.append(cells)
            continue
        if table_rows: add_table(doc, table_rows); table_rows=[]
        if not text or text == '---': continue
        if text.startswith('# '):
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(4); add_text(p, text[2:], 24, RGBColor(25,31,40), True)
        elif text.startswith('> '):
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(16); add_text(p, text[2:], 10, RGBColor(107,118,132))
        elif text.startswith('## '):
            p=doc.add_paragraph(style='Heading 1'); add_text(p, text[3:], 16, BLUE, True)
        elif text.startswith('### '):
            p=doc.add_paragraph(style='Heading 2'); add_text(p, text[4:], 13, BLUE, True)
        elif text.startswith('- [ ] '):
            p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(4); add_text(p, '☐ ' + text[6:], 11)
        elif text.startswith('- '):
            p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(4); add_text(p, text[2:], 11)
        elif re.match(r'^\d+\. ', text):
            p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(4); add_text(p, re.sub(r'^\d+\. ', '', text), 11)
        else:
            p=doc.add_paragraph(); add_text(p, text, 11)
    if table_rows: add_table(doc, table_rows)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True); doc.save(OUTPUT)

if __name__ == '__main__': build()
