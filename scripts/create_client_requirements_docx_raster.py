from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from pypdfium2 import PdfDocument

ROOT = Path(__file__).resolve().parents[1]
PDF = ROOT / 'output/pdf/timefit-client-requirements-detailed-review.pdf'
OUTPUT = ROOT / 'output/docx/timefit-client-requirements-notion-guide.docx'
TMP = ROOT / 'tmp/docx-client-requirements-source'

def build():
    TMP.mkdir(parents=True, exist_ok=True)
    pdf = PdfDocument(str(PDF))
    images = []
    for index in range(len(pdf)):
        image_path = TMP / f'page-{index + 1}.png'
        pdf[index].render(scale=2).to_pil().save(image_path)
        images.append(image_path)
    doc = Document(); section = doc.sections[0]
    section.top_margin = Inches(0.5); section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.8); section.right_margin = Inches(0.8)
    for index, image_path in enumerate(images):
        paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Inches(0); paragraph.paragraph_format.space_after = Inches(0)
        paragraph.add_run().add_picture(str(image_path), width=Inches(6.9))
        if index < len(images) - 1: doc.add_page_break()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True); doc.save(OUTPUT)

if __name__ == '__main__': build()
